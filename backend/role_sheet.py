# -*- coding: utf-8 -*-
"""角色表格式转换：
{Gender, Name(zh), Name(en), Description(zh)} → {Gender, Source, Target, Intro, Profile}
每条数据行后追加 N 行空行；支持 xlsx 与 csv 输入。
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    load_workbook = None
    Workbook = None


# 列名候选：按优先级匹配，命中任意一个就算该列
COLUMN_MAP = {
    'gender':  ['gender', '性别'],
    'name_zh': ['name (zh)', 'name(zh)', 'name_zh', '中文名', '中文姓名'],
    'name_en': ['name (en)', 'name(en)', 'name_en', '英文名', '英文姓名', 'english name'],
    'desc_zh': ['description (zh)', 'description(zh)', 'description', '简介', '描述', '备注', 'profile'],
}

NEW_HEADERS = ['Gender', 'Source', 'Target', 'Intro', 'Profile']
DEFAULT_SHEET_TITLE = '本地化清单'


def _norm(s):
    return str(s or '').strip().lower()


# ----------------------------------------------------------------
# 读取层：xlsx / csv → 统一的 rows[[v, v, ...]] + header 位置字典
# ----------------------------------------------------------------

def _read_xlsx(src_bytes: bytes):
    if load_workbook is None:
        raise RuntimeError("缺少依赖 openpyxl：pip install openpyxl")
    wb = load_workbook(io.BytesIO(src_bytes))
    ws = wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
    return rows


def _read_csv(src_bytes: bytes):
    # CSV 可能有 BOM（UTF-8-SIG）或 GBK；多种编码尝试
    text = None
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030'):
        try:
            text = src_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 无法识别编码，请另存为 UTF-8 或 GBK")

    # 尝试分号分隔（某些欧洲/本地化工具默认用分号）
    sample = text.splitlines()[:3]
    dialect = None
    try:
        dialect = csv.Sniffer().sniff('\n'.join(sample), delimiters='\t,;')
    except csv.Error:
        dialect = csv.excel  # 默认逗号

    rows = []
    reader = csv.reader(text.splitlines(), dialect)
    for row in reader:
        rows.append([(c if c is not None else '') for c in row])
    return rows


def _detect_format(filename: str):
    """按扩展名判断。"""
    name = (filename or '').lower()
    if name.endswith('.csv'):
        return 'csv'
    return 'xlsx'  # 默认


def _read_source(src_bytes: bytes, filename: str):
    fmt = _detect_format(filename)
    if fmt == 'csv':
        return _read_csv(src_bytes)
    return _read_xlsx(src_bytes)


# ----------------------------------------------------------------
# 解析层：rows → 列索引字典
# ----------------------------------------------------------------

def _find_columns(rows):
    """扫首行（跳过空表头行）返回字段 key → 列索引。"""
    header_row = None
    for r in rows:
        if any(str(c).strip() for c in r):
            header_row = r
            break
    if header_row is None:
        return {}

    header = {}
    for col, v in enumerate(header_row):
        v_n = _norm(v)
        if not v_n:
            continue
        for key, candidates in COLUMN_MAP.items():
            if key in header:
                continue
            for cand in candidates:
                cand_n = _norm(cand)
                if cand_n == v_n or cand_n in v_n:
                    header[key] = col
                    break
    return header


# ----------------------------------------------------------------
# 主入口
# ----------------------------------------------------------------

def convert(
    src_bytes: bytes,
    filename: str = 'role_sheet.xlsx',
    blank_rows: int = 5,
    sheet_name: str = DEFAULT_SHEET_TITLE,
    preview_only: bool = False,
    preview_rows: int = 3,
):
    """转换源文件。

    preview_only=True 时只返回预览数据（前 preview_rows 条角色），不生成 xlsx。
    否则返回 (out_xlsx_bytes, report_dict)。
    """
    rows = _read_source(src_bytes, filename)
    cols = _find_columns(rows)

    missing = [k for k in ('name_zh', 'desc_zh') if k not in cols]
    if missing:
        raise ValueError(f"表头缺少必要列: {missing}（识别到的字段列: {cols}）")

    def _get(row, key):
        idx = cols.get(key)
        if idx is None or idx >= len(row):
            return ''
        return row[idx] if row[idx] is not None else ''

    # 收集全部转换后的角色数据（跳过空行）
    converted = []
    skipped = 0
    for row in rows[1:]:
        if not row or all((c is None or str(c).strip() == '') for c in row):
            skipped += 1
            continue
        converted.append({
            'gender':  _get(row, 'gender'),
            'source':  _get(row, 'name_zh'),
            'target':  '',
            'intro':   '',
            'profile': _get(row, 'desc_zh'),
        })

    report = {
        'written': len(converted),
        'skipped_blank_rows': skipped,
        'blank_rows_between': blank_rows,
        'columns_detected': cols,
        'sheet_name': sheet_name,
        'source_format': _detect_format(filename),
    }

    if preview_only:
        report['preview'] = converted[:preview_rows]
        return None, report

    # 写输出 xlsx
    if Workbook is None:
        raise RuntimeError("缺少依赖 openpyxl：pip install openpyxl")

    out = Workbook()
    ows = out.active
    ows.title = sheet_name
    ows.append(NEW_HEADERS)

    for item in converted:
        ows.append([item['gender'], item['source'], '', '', item['profile']])
        for _ in range(blank_rows):
            ows.append(['', '', '', '', ''])

    out_buf = io.BytesIO()
    out.save(out_buf)
    out_buf.seek(0)
    out_bytes = out_buf.getvalue()
    return out_bytes, report
